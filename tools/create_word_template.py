from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from uuid import uuid4
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
import xml.etree.ElementTree as ET

out = Path(__file__).resolve().parents[1] / 'word'
out.mkdir(parents=True, exist_ok=True)
store = '{' + str(uuid4()).upper() + '}'
doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)
doc.add_heading('Produktbeleg', 0)
doc.add_paragraph('Ein Dokument pro Datensatz', style='Subtitle')
table = doc.add_table(rows=0, cols=2)
table.style = 'Light Shading Accent 1'
fields = [('ID', 'id', '1001'), ('Vorname', 'vorname', 'Anna'), ('Nachname', 'nachname', 'Müller'), ('Produkt', 'produkt', 'Notizbuch'), ('Preis (EUR)', 'preis', '12.90'), ('Datum', 'datum', '2026-10-11')]
for index, (label, key, value) in enumerate(fields, 1):
    cells = table.add_row().cells
    cells[0].text = label
    p = cells[1].paragraphs[0]
    sdt = OxmlElement('w:sdt')
    props = OxmlElement('w:sdtPr')
    for tag, val in [('alias', label), ('tag', key), ('id', str(1000+index))]:
        el = OxmlElement('w:'+tag); el.set(qn('w:val'), val); props.append(el)
    props.append(OxmlElement('w:text'))
    binding = OxmlElement('w:dataBinding')
    binding.set(qn('w:storeItemID'), store)
    binding.set(qn('w:xpath'), '/datensaetze/datensatz[1]/' + ('@id' if key == 'id' else key))
    props.append(binding)
    sdt.append(props)
    content = OxmlElement('w:sdtContent')
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = value
    run.append(text); content.append(run); sdt.append(content); p._p.append(sdt)
doc.add_paragraph('Preis in EUR. Datum im Format JJJJ-MM-TT.')
buf = BytesIO(); doc.save(buf)
xml = '''<?xml version="1.0" encoding="utf-8"?>
<datensaetze version="1"><datensatz id="1001"><vorname>Anna</vorname><nachname>Müller</nachname><produkt>Notizbuch</produkt><preis waehrung="EUR">12.90</preis><datum>2026-10-11</datum></datensatz></datensaetze>'''
rel_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
ct_ns = 'http://schemas.openxmlformats.org/package/2006/content-types'
with ZipFile(buf) as source:
    files = {name: source.read(name) for name in source.namelist()}
rels = ET.fromstring(files['word/_rels/document.xml.rels'])
# python-docx already includes this custom XML relationship.
if not any(el.get('Target') == '../customXml/item1.xml' for el in rels):
    ET.SubElement(rels, '{'+rel_ns+'}Relationship', Id='rIdXMLLearning', Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml', Target='../customXml/item1.xml')
files['word/_rels/document.xml.rels'] = ET.tostring(rels, encoding='utf-8', xml_declaration=True)
files['customXml/item1.xml'] = xml.encode()
files['customXml/itemProps1.xml'] = f'''<?xml version="1.0"?><ds:datastoreItem ds:itemID="{store}" xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml"><ds:schemaRefs/></ds:datastoreItem>'''.encode()
files['customXml/_rels/item1.xml.rels'] = f'''<?xml version="1.0"?><Relationships xmlns="{rel_ns}"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps" Target="itemProps1.xml"/></Relationships>'''.encode()
cts = ET.fromstring(files['[Content_Types].xml'])
# Keep exactly one content-type override per package part.
if not any(el.get('PartName') == '/customXml/itemProps1.xml' for el in cts):
    ET.SubElement(cts, '{'+ct_ns+'}Override', PartName='/customXml/itemProps1.xml', ContentType='application/vnd.openxmlformats-officedocument.customXmlProperties+xml')
for filename, template in [('Produktbeleg-Beispiel.docx', False), ('Produktbeleg-Vorlage.dotx', True)]:
    for el in cts:
        if el.attrib.get('PartName') == '/word/document.xml':
            el.set('ContentType', 'application/vnd.openxmlformats-officedocument.wordprocessingml.' + ('template' if template else 'document') + '.main+xml')
    files['[Content_Types].xml'] = ET.tostring(cts, encoding='utf-8', xml_declaration=True)
    with ZipFile(out / filename, 'w', ZIP_DEFLATED) as archive:
        for name, data in files.items(): archive.writestr(name, data)
(out / 'Beispieldaten.xml').write_text(xml, encoding='utf-8')
print('Word-Vorlage, Beispieldokument und Beispiel-XML erstellt.')
